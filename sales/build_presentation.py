"""Elecnovatech HYBRID sales presentation.

Claims in this deck are limited to behavior present in the product.
Bracketed fields are filled in for each client. Nothing here is a price.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from pathlib import Path

ROOT = Path(__file__).resolve().parent
MARK = ROOT / "assets" / "hybrid-mark.png"
OUT = ROOT / "Elecnovatech-HYBRID-Sales-Presentation.pptx"

NAVY = RGBColor(0x1A, 0x20, 0x2C)
INK = RGBColor(0x1A, 0x20, 0x2C)
DEEP = RGBColor(0x0E, 0x1A, 0x26)
GREEN = RGBColor(0x80, 0xC3, 0x42)
GREEN_DARK = RGBColor(0x2D, 0x5A, 0x27)
BLUE = RGBColor(0x00, 0xAE, 0xEF)
BLUE_DARK = RGBColor(0x00, 0x54, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF6, 0xF8, 0xF4)
MUTED = RGBColor(0x5C, 0x6B, 0x73)
LINE = RGBColor(0xE1, 0xE8, 0xDF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_GREEN = RGBColor(0xF3, 0xF8, 0xEC)
SOFT_BLUE = RGBColor(0xF0, 0xF8, 0xFC)
AMBER_BG = RGBColor(0xFB, 0xF6, 0xEA)
SLATE = RGBColor(0x3D, 0x4A, 0x57)

W = 13.333
H = 7.5
TOTAL = 16


def _set_run(run, text, size, bold, color, font="Calibri"):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("a:latin"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("a:latin"), {})
        rpr.append(rfonts)
    rfonts.set("typeface", font)


def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        _set_run(run, line, size, bold, color, font)
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if radius:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    sppr = shape._element.spPr
    effect = sppr.find(qn("a:effectLst"))
    if effect is not None:
        sppr.remove(effect)
    return shape


def bg(slide, color):
    rect(slide, 0, 0, W, H, color)


def footer(slide, number, dark=False):
    color = RGBColor(0x9A, 0xA8, 0xB0) if dark else MUTED
    rule = RGBColor(0x2A, 0x3A, 0x48) if dark else LINE
    rect(slide, 0.5, 7.12, 12.33, 0.01, rule)
    add_text(slide, "ELECNOVATECH", 0.5, 7.18, 2.2, 0.26, 10, True, color)
    add_text(slide, "HYBRID  ·  Sales presentation", 2.7, 7.18, 6.5, 0.26, 10, False, color, PP_ALIGN.CENTER)
    add_text(slide, f"{number:02d}  /  {TOTAL:02d}", 10.4, 7.18, 2.43, 0.26, 10, True, color, PP_ALIGN.RIGHT)


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, DEEP if dark else CREAM)
    return slide


def kicker(slide, text, x, y, w=6):
    add_text(slide, text.upper(), x, y, w, 0.24, 11, True, GREEN_DARK)


def title(slide, text, x, y, w, h=0.55, size=30, color=NAVY):
    add_text(slide, text, x, y, w, h, size, True, color)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def card(slide, x, y, w, h, fill=CARD):
    return rect(slide, x, y, w, h, fill, LINE, 0.08)


def accent_bar(slide, x, y, w, color):
    rect(slide, x, y, w, 0.06, color)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.author = "Elecnovatech"
    prs.core_properties.title = "HYBRID — Sales presentation"
    prs.core_properties.subject = "Operations platform for block factories and plastic recycling plants"
    prs.core_properties.category = "Sales"
    prs.core_properties.comments = (
        "Editable sales deck. Bracketed fields are completed per client. "
        "Do not add metrics, testimonials, or features that are not in the product."
    )

    cover(prs)
    problem(prs)
    audience(prs)
    benefits(prs)
    modules(prs)
    blocks(prs)
    recycling(prs)
    sales(prs)
    finance(prs)
    control(prs)
    floor(prs)
    scope(prs)
    implementation(prs)
    commercial(prs)
    why(prs)
    close(prs)

    prs.save(OUT)
    print(OUT)


def cover(prs):
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, 0.14, H, GREEN)
    rect(s, 0.14, 0, 0.06, H, BLUE)
    s.shapes.add_picture(str(MARK), Inches(0.7), Inches(0.42), Inches(1.55), Inches(1.16))
    add_text(s, "ELECNOVATECH", 2.4, 0.62, 5, 0.28, 13, True, GREEN)
    add_text(s, "Product introduction", 2.4, 0.92, 5, 0.28, 14, False, BLUE)
    add_text(s, "HYBRID", 0.7, 1.85, 10, 0.95, 60, True, WHITE)
    add_text(
        s,
        "The operations platform for block factories\nand plastic recycling plants.",
        0.72,
        2.9,
        9,
        0.9,
        24,
        False,
        RGBColor(0xD5, 0xE4, 0xC8),
    )
    rect(s, 0.72, 4.05, 2.1, 0.06, GREEN)
    add_text(
        s,
        "Developed on Elecnovatech’s own block plant and recycling plant.\nConfigured and licensed again for each new client.",
        0.72,
        4.28,
        8.5,
        0.7,
        16,
        False,
        RGBColor(0xC5, 0xD0, 0xD6),
    )

    pills = [
        (0.72, "Block production"),
        (3.35, "Plastic recycling"),
        (6.15, "Ethiopian Birr"),
        (8.7, "One client at a time"),
    ]
    for x, label in pills:
        rect(s, x, 5.3, 2.45, 0.46, RGBColor(0x16, 0x28, 0x36), radius=0.3)
        add_text(s, label, x, 5.38, 2.45, 0.32, 13, True, WHITE, PP_ALIGN.CENTER)

    add_text(s, "Prepared for  [Client name]", 0.72, 6.15, 5.5, 0.28, 14, False, RGBColor(0x9A, 0xB4, 0xA0))
    add_text(s, "[Presentation date]", 8.2, 6.15, 4.4, 0.28, 14, False, RGBColor(0x9A, 0xB4, 0xA0), PP_ALIGN.RIGHT)
    footer(s, 1, dark=True)
    notes(
        s,
        "Open by naming the client’s sites. HYBRID is the product. Elecnovatech is the company. "
        "The first installation is Elecnovatech’s own block factory and plastic recycling factory. "
        "Replace the bracketed client name and date before the meeting.",
    )


def problem(prs):
    s = new_slide(prs)
    kicker(s, "The situation", 0.5, 0.28)
    title(s, "The factory day is split across too many records.", 0.5, 0.52, 12)
    items = [
        ("01", "Production", "Batch output, damage, and material use have to travel from the machine to someone who can trust the number."),
        ("02", "Stock", "Cement, poliche, finished blocks, plastic input, and finished pellets each need a balance the yard and the office share."),
        ("03", "Credit", "A load can leave before finance can see what that customer still owes, including withholding and earlier payments."),
        ("04", "The books", "Many plants already keep accounts in Sage 50. The factory record still has to be compared with those books."),
    ]
    positions = [(0.5, 1.55), (6.9, 1.55), (0.5, 4.15), (6.9, 4.15)]
    for (num, head, body), (x, y) in zip(items, positions):
        card(s, x, y, 5.95, 2.35)
        accent_bar(s, x, y, 5.95, GREEN if num in ("01", "03") else BLUE)
        add_text(s, num, x + 0.28, y + 0.28, 1.1, 0.4, 20, True, GREEN_DARK)
        add_text(s, head, x + 1.3, y + 0.3, 4.2, 0.38, 20, True, NAVY)
        add_text(s, body, x + 0.28, y + 0.95, 5.4, 1.15, 15, False, SLATE)
    footer(s, 2)
    notes(s, "Stay with the client’s own examples. Do not quote a percentage of waste, leakage, or time saved. We have not measured that for them.")


def audience(prs):
    s = new_slide(prs)
    kicker(s, "Who it is for", 0.5, 0.28)
    title(s, "A fit when the work looks like this.", 0.5, 0.52, 12)
    rows = [
        ("Block plant", "You make concrete blocks, track cement and poliche, and sell finished sizes."),
        ("Recycling plant", "You receive plastic, record processed output and waste in kilograms, and sell the result."),
        ("Both sites", "You want one set of customers, suppliers, banks, and books, with a separate view for each site."),
        ("Real job split", "Owner, finance, factory, warehouse, sales, HR, floor, and driver do not all need the same screen."),
    ]
    y = 1.5
    for i, (head, body) in enumerate(rows):
        card(s, 0.5, y, 12.33, 1.2)
        rect(s, 0.5, y, 0.1, 1.2, GREEN if i % 2 == 0 else BLUE)
        add_text(s, head, 0.9, y + 0.18, 3.1, 0.84, 18, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, body, 4.1, y + 0.18, 8.4, 0.84, 16, False, SLATE, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.32
    footer(s, 3)
    notes(
        s,
        "A client with only one of the two sites can work in that view. "
        "Renaming the sites, adding a third plant, or covering a different industry is quoted as customization. "
        "Each client gets their own deployment. Unrelated companies do not share one login portal.",
    )


def benefits(prs):
    s = new_slide(prs)
    kicker(s, "What changes at work", 0.5, 0.28)
    title(s, "Each role keeps the part of the day it owns.", 0.5, 0.52, 12)
    cards = [
        (GREEN, "Owner", "Switch between the company, the block plant, and the recycling plant. See orders, credit, and output without collecting separate files."),
        (BLUE, "Finance", "Statements, aging, payments, expenses, and a workspace that compares factory amounts with Sage 50."),
        (GREEN_DARK, "Factory and warehouse", "Batches, daily targets, stock balances, physical counts, and how far current cement and poliche can go."),
        (BLUE_DARK, "Floor, drivers, HR", "Floor staff and drivers get a short phone screen. HR runs attendance, leave, and payroll from rules the company sets."),
    ]
    positions = [(0.5, 1.5), (6.9, 1.5), (0.5, 4.15), (6.9, 4.15)]
    for (color, head, body), (x, y) in zip(cards, positions):
        card(s, x, y, 5.95, 2.4)
        rect(s, x, y + 0.28, 0.08, 0.42, color)
        add_text(s, head, x + 0.32, y + 0.26, 5.3, 0.42, 20, True, NAVY)
        add_text(s, body, x + 0.32, y + 0.85, 5.3, 1.3, 15, False, SLATE)
    footer(s, 4)
    notes(s, "Salary visibility is a separate permission. Do not imply every user can see payroll.")


def modules(prs):
    s = new_slide(prs)
    kicker(s, "The working surface", 0.5, 0.28)
    title(s, "One system. The screens follow the job.", 0.5, 0.52, 12)
    tiles = [
        ("Dashboard", "Company or one site"),
        ("Production", "Block batches and counts"),
        ("Recycling", "Kilograms in and out"),
        ("Inventory", "Balances and counts"),
        ("Purchasing", "Suppliers and receipts"),
        ("Sales", "Orders and customers"),
        ("Credit", "Who still owes"),
        ("Finance", "Books and Sage compare"),
        ("Banks", "Collections by account"),
        ("Expenses", "Operating costs"),
        ("HR", "People, leave, pay"),
        ("Reports", "PDF and Excel"),
    ]
    for i, (head, body) in enumerate(tiles):
        col = i % 4
        row = i // 4
        x = 0.5 + col * 3.2
        y = 1.5 + row * 1.75
        card(s, x, y, 3.02, 1.58)
        rect(s, x + 0.2, y + 0.22, 0.42, 0.1, GREEN if row != 1 else BLUE)
        add_text(s, head, x + 0.2, y + 0.42, 2.62, 0.4, 16, True, NAVY)
        add_text(s, body, x + 0.2, y + 0.88, 2.62, 0.48, 13, False, MUTED)
    footer(s, 5)
    notes(
        s,
        "Also in the product, and not on this grid: documents, notifications, audit log, admin roles, and the phone screens. "
        "The movements list and the bill-of-materials screen exist in the code and are switched off in the current inventory page. "
        "Treat those as a proposed switch-on, not as a live menu item.",
    )


def blocks(prs):
    s = new_slide(prs)
    kicker(s, "Block factory", 0.5, 0.28)
    title(s, "From a batch to a size you can sell.", 0.5, 0.52, 10)
    points = [
        "Batches for 10, 15, 20, and 24 cm, with formulas for cement bags, plates, poliche loads, and pallets.",
        "Counts at three points: batch count, after drying, and after curing, plus daily output.",
        "Yield and damage calculated from the counts you enter.",
        "Daily targets, and a capacity view from the cement and poliche on hand.",
        "An in-app calculator for block quantities, waste, and material.",
    ]
    y = 1.45
    for i, point in enumerate(points, 1):
        rect(s, 0.5, y, 0.42, 0.42, GREEN if i % 2 else BLUE_DARK, radius=0.2)
        add_text(s, str(i), 0.5, y + 0.04, 0.42, 0.34, 14, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, point, 1.12, y, 7.15, 0.7, 15, False, SLATE)
        y += 0.95

    card(s, 8.55, 1.5, 4.28, 5.15, SOFT_GREEN)
    add_text(s, "SIZES IN THE PRODUCT", 8.8, 1.72, 3.85, 0.28, 12, True, GREEN_DARK)
    for i, size in enumerate(["10 cm", "15 cm", "20 cm", "24 cm"]):
        yy = 2.2 + i * 0.85
        rect(s, 8.85, yy, 3.7, 0.68, WHITE, LINE, 0.12)
        add_text(s, size, 8.85, yy + 0.14, 3.7, 0.4, 18, True, NAVY, PP_ALIGN.CENTER)
    add_text(s, "Formulas and bag weight are settings.", 8.8, 5.7, 3.85, 0.6, 13, False, MUTED)
    footer(s, 6)
    notes(
        s,
        "Default formulas ship with the product and can be changed in admin settings. "
        "Confirm the client’s real sizes and bag weight in discovery. "
        "Capacity names the limiting material. Confidence rises as daily history builds. Do not promise a forecast accuracy rate.",
    )


def recycling(prs):
    s = new_slide(prs)
    kicker(s, "Plastic recycling", 0.5, 0.28)
    title(s, "Kilograms in. Kilograms out. Waste in between.", 0.5, 0.52, 12)
    steps = [
        ("1", "Received", "Material in, by kilogram"),
        ("2", "Processed", "Output recorded"),
        ("3", "Waste", "The gap is the waste"),
        ("4", "Stock", "Raw and finished balances"),
        ("5", "Order", "A stage on the order"),
    ]
    for i, (n, head, body) in enumerate(steps):
        x = 0.5 + i * 2.55
        card(s, x, 1.55, 2.4, 2.15)
        add_text(s, n, x + 0.16, 1.7, 0.4, 0.36, 16, True, BLUE_DARK)
        add_text(s, head, x + 0.16, 2.15, 2.08, 0.4, 16, True, NAVY)
        add_text(s, body, x + 0.16, 2.6, 2.08, 0.8, 13, False, SLATE)
    card(s, 0.5, 3.95, 12.33, 2.85)
    add_text(s, "What the recycling screens already do", 0.78, 4.15, 8, 0.36, 16, True, NAVY)
    cols = [
        "Today, this week, and this month output on the recycling view.",
        "Yield from received kilograms and processed kilograms.",
        "Order stages: pending, accepted, in production, ready, delivered, cancelled.",
        "Planning uses those records. It does not reuse the block-size learning model.",
    ]
    for i, line in enumerate(cols):
        col = i % 2
        row = i // 2
        add_text(s, "●   " + line, 0.78 + col * 6.05, 4.7 + row * 0.85, 5.8, 0.75, 14, False, SLATE)
    footer(s, 7)
    notes(s, "Recycling orders keep a stage in addition to the simple pending-pickup and picked status used for blocks.")


def sales(prs):
    s = new_slide(prs)
    kicker(s, "Sales, credit, delivery", 0.5, 0.28)
    title(s, "A sale stays attached to the customer who owes it.", 0.5, 0.52, 12)
    left = [
        ("Customers", "Shared across both sites. Regular or walking. TIN can be stored."),
        ("Orders", "Pending pickup, picked, or cancelled. VAT, discount, withholding, and advances sit on the order."),
        ("Credit", "The balance already reflects withholding and store credit. A score from 0 to 100 is calculated from that customer’s own history, per site and combined."),
        ("Proof", "The driver closes delivery with a photo, a signature, and a GPS point."),
    ]
    y = 1.48
    for head, body in left:
        add_text(s, head, 0.5, y, 2.3, 0.36, 16, True, GREEN_DARK)
        add_text(s, body, 2.9, y, 6.3, 1.05, 14, False, SLATE)
        y += 1.25
    card(s, 9.4, 1.5, 3.43, 5.2, DEEP)
    add_text(s, "ON THE PHONE", 9.65, 1.75, 3.0, 0.3, 12, True, GREEN)
    add_text(s, "Delivery\nconfirmation", 9.65, 2.25, 3.0, 0.9, 22, True, WHITE)
    add_text(s, "Photo", 9.65, 3.4, 3.0, 0.35, 15, False, RGBColor(0xD0, 0xDE, 0xC8))
    add_text(s, "Signature", 9.65, 3.85, 3.0, 0.35, 15, False, RGBColor(0xD0, 0xDE, 0xC8))
    add_text(s, "GPS point", 9.65, 4.3, 3.0, 0.35, 15, False, RGBColor(0xD0, 0xDE, 0xC8))
    add_text(s, "Queued if the signal drops, then sent when it returns.", 9.65, 5.15, 2.95, 1.1, 13, False, RGBColor(0xB7, 0xC4, 0xCC))
    footer(s, 8)
    notes(
        s,
        "Payment methods on an order are cash or bank. Invoice required is a flag on the order. "
        "Do not describe a card machine, Telebirr, or an e-invoice integration. Those are not in the product. "
        "Credit scores are calculated in the application. They are not a bank decision.",
    )


def finance(prs):
    s = new_slide(prs)
    kicker(s, "Finance", 0.5, 0.28)
    title(s, "The factory record and Sage 50 can be compared.", 0.5, 0.52, 12)
    tiles = [
        ("Period view", "Revenue, collections, receivables, and cash signals for a chosen period."),
        ("Customer money", "Statements, aging buckets, and a payment report."),
        ("Cost and banks", "Expenses, profitability, bank accounts, and bank transactions."),
        ("Tax settings", "VAT and withholding are settings. They apply on new orders."),
    ]
    for i, (head, body) in enumerate(tiles):
        col = i % 2
        row = i // 2
        x = 0.5 + col * 6.4
        y = 1.48 + row * 1.7
        card(s, x, y, 6.15, 1.55)
        add_text(s, head, x + 0.25, y + 0.18, 5.7, 0.36, 16, True, NAVY)
        add_text(s, body, x + 0.25, y + 0.62, 5.7, 0.72, 14, False, SLATE)
    rect(s, 0.5, 5.0, 12.33, 1.85, SOFT_BLUE, radius=0.06)
    add_text(s, "SAGE 50  ·  CONFIRMED PATH, CLIENT-SPECIFIC CONNECTOR", 0.75, 5.18, 11.8, 0.3, 12, True, BLUE_DARK)
    add_text(
        s,
        "A Windows agent on the Sage PC sends data one way into HYBRID, on a timer or with Sync Now. The copy that ships reads an export file. During implementation we complete the reader for that client’s Sage 50 version. HYBRID does not replace Sage.",
        0.75,
        5.55,
        11.8,
        1.05,
        14,
        False,
        SLATE,
    )
    footer(s, 9)
    notes(
        s,
        "Say this sentence out loud: the agent is not a finished connector for every Sage 50 version. "
        "Quote the connector after you know the version. The link is outbound from their PC. "
        "There is a reconciliation status of matched, variance, pending, or approved.",
    )


def control(prs):
    s = new_slide(prs)
    kicker(s, "Control", 0.5, 0.28)
    title(s, "Stock, people, and a record of who changed what.", 0.5, 0.52, 12)
    columns = [
        (GREEN, "Inventory", [
            "Live balances and a stock ledger",
            "Physical counts, including from the phone",
            "Low-stock markers",
            "Catalog of types, categories, and units",
            "Purchasing and supplier balances",
        ]),
        (BLUE, "People", [
            "Employee records and attendance",
            "Leave types, holidays, and approvals",
            "Payroll calculated from the company’s rates, shifts, and attendance",
            "Payslips, with salary hidden unless the role allows it",
            "Attendance corrections",
        ]),
        (GREEN_DARK, "Governance", [
            "Roles from owner to driver",
            "Permissions stored as data and editable",
            "Audit log on key records",
            "Documents, with tags and restricted files",
            "In-app notifications",
        ]),
    ]
    for i, (color, head, lines) in enumerate(columns):
        x = 0.5 + i * 4.25
        card(s, x, 1.48, 4.05, 5.25)
        rect(s, x, 1.48, 4.05, 0.7, color)
        add_text(s, head, x + 0.2, 1.6, 3.65, 0.46, 18, True, WHITE)
        yy = 2.4
        for line in lines:
            add_text(s, line, x + 0.22, yy, 3.6, 0.7, 14, False, SLATE)
            yy += 0.78
    footer(s, 10)
    notes(
        s,
        "Payroll calculates amounts inside HYBRID. It does not file tax returns. "
        "Bill of materials, locations, and lots are built and currently hidden. Offer them as a switch-on during scoping. "
        "Report schedules can be stored and paused. Automatic email delivery is not a verified mail integration.",
    )


def floor(prs):
    s = new_slide(prs)
    kicker(s, "Floor tools and planning", 0.5, 0.28)
    title(s, "Short actions on the phone. Numbers from the records.", 0.5, 0.52, 12)
    card(s, 0.5, 1.5, 6.15, 5.25)
    add_text(s, "PHONE  ·  CONFIRMED", 0.75, 1.7, 5.6, 0.3, 12, True, GREEN_DARK)
    add_text(s, "A phone browser, organized by role.", 0.75, 2.08, 5.6, 0.4, 18, True, NAVY)
    phone = [
        "Floor — production count and stock count",
        "Warehouse — stock count and a movement",
        "Driver — delivery proof only",
        "Factory manager — all four actions",
        "Works from a queue when the phone is offline",
    ]
    yy = 2.7
    for line in phone:
        add_text(s, line, 0.75, yy, 5.6, 0.55, 15, False, SLATE)
        yy += 0.68

    card(s, 6.9, 1.5, 5.93, 5.25, SOFT_GREEN)
    add_text(s, "PLANNING  ·  CONFIRMED", 7.15, 1.7, 5.45, 0.3, 12, True, GREEN_DARK)
    add_text(s, "Calculated in the product.", 7.15, 2.08, 5.45, 0.4, 18, True, NAVY)
    plan = [
        "Block capacity from cement and poliche on hand, by size.",
        "Waste and yield drawn from daily history, with a confidence level.",
        "Customer credit score, tier, and limit from that customer’s payments.",
        "Plain-language answers only when an AI key is configured. The figures still come from HYBRID.",
    ]
    yy = 2.7
    for line in plan:
        add_text(s, line, 7.15, yy, 5.4, 0.85, 14, False, SLATE)
        yy += 0.9
    footer(s, 11)
    notes(
        s,
        "Do not call this a Play Store or App Store app. It is the same web system on a phone. "
        "Do not promise answers if the client will not provide an xAI or OpenAI key. "
        "The design rule in the product is that numbers are computed in code.",
    )


def scope(prs):
    s = new_slide(prs)
    kicker(s, "How to talk about scope", 0.5, 0.28)
    title(s, "Three words we keep separate in every sale.", 0.5, 0.52, 12)
    cols = [
        (GREEN_DARK, WHITE, "Confirmed", "In the product today", [
            "Two site views plus a company view",
            "The modules on the previous slides",
            "Birr, VAT, withholding, TIN",
            "PDF and Excel export",
            "Role permissions and audit",
            "Phone actions and offline queue",
        ]),
        (BLUE_DARK, WHITE, "Configurable", "Settings, not new code", [
            "Yield formulas and bag weight",
            "Daily targets",
            "Tax rates and HR rules",
            "Catalog, banks, and roles",
            "Which site a person works in",
            "Opening balances at go-live",
        ]),
        (RGBColor(0x8A, 0x6A, 0x1A), INK, "Proposed", "Quoted before it is promised", [
            "Sage reader for their version",
            "Amharic screens",
            "Hidden stock tools switched on",
            "A third site, or a new process",
            "Emailed report schedules",
            "Support terms and the fee",
        ]),
    ]
    for i, (header, header_ink, head, sub, lines) in enumerate(cols):
        x = 0.5 + i * 4.25
        card(s, x, 1.48, 4.05, 5.25)
        rect(s, x, 1.48, 4.05, 1.15, header)
        add_text(s, head, x + 0.2, 1.55, 3.65, 0.42, 20, True, header_ink if i < 2 else WHITE)
        add_text(s, sub, x + 0.2, 2.02, 3.65, 0.4, 12, False, RGBColor(0xE5, 0xF2, 0xD8) if i < 2 else RGBColor(0xFF, 0xF4, 0xD8))
        yy = 2.85
        for line in lines:
            add_text(s, line, x + 0.22, yy, 3.6, 0.55, 14, False, SLATE)
            yy += 0.58
    footer(s, 12)
    notes(
        s,
        "If a buyer asks for a feature in the right column, write it into the scope and the price. "
        "The amber column header uses white text. Keep the three labels in the proposal as well.",
    )


def implementation(prs):
    s = new_slide(prs)
    kicker(s, "Implementation", 0.5, 0.28)
    title(s, "A separate deployment for that client.", 0.5, 0.52, 12)
    steps = [
        ("1", "Discover", "Sites, sizes, taxes, roles, banks, and whether Sage 50 is in use."),
        ("2", "Deploy", "Their own database and sign-in. One organization."),
        ("3", "Configure", "Yields, catalog, tax, HR rules, roles, and banks."),
        ("4", "Open", "Opening stock, customers, suppliers, and employees."),
        ("5", "Connect", "Sage agent, after the version is known. Skip if they do not use Sage."),
        ("6", "Train", "A session per role, then go-live on an agreed date."),
    ]
    for i, (n, head, body) in enumerate(steps):
        col = i % 3
        row = i // 3
        x = 0.5 + col * 4.25
        y = 1.55 + row * 2.45
        card(s, x, y, 4.05, 2.25)
        rect(s, x + 0.22, y + 0.28, 0.48, 0.48, GREEN if row == 0 else BLUE_DARK, radius=0.2)
        add_text(s, n, x + 0.22, y + 0.34, 0.48, 0.36, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, head, x + 0.86, y + 0.32, 2.9, 0.42, 18, True, NAVY)
        add_text(s, body, x + 0.22, y + 1.0, 3.6, 0.95, 14, False, SLATE)
    footer(s, 13)
    notes(
        s,
        "Timetable is [Agreed timetable] in the quote. Do not promise a fixed number of weeks in the room. "
        "Opening data is entered during setup. Do not promise a spreadsheet import unless it is added to the scope.",
    )


def commercial(prs):
    s = new_slide(prs)
    kicker(s, "Commercial structure", 0.5, 0.26)
    title(s, "The quote fills every bracket.", 0.5, 0.48, 12, size=28)
    rows = [
        ("Item", "What the client is buying", "Amount"),
        ("License", "Use of HYBRID for one organization, for the term in the quote", "[Annual fee]"),
        ("Implementation", "Discovery, deployment, configuration, opening data, training, go-live", "[Project fee]"),
        ("Sage connector", "Reader for their Sage 50 version, and the agent on their PC", "[Fee or not required]"),
        ("Customization", "Only the written list. Anything else stays out of this sale", "[Scoped fee]"),
        ("Support", "The retainer described in the proposal, after go-live", "[Monthly fee]"),
    ]
    table_shape = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.4), Inches(12.33), Inches(4.55))
    table = table_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(7.15)
    table.columns[2].width = Inches(2.68)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.08)
            cell.margin_bottom = Inches(0.06)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            is_header = r == 0
            is_amount = c == 2 and r > 0
            _set_run(
                run,
                value,
                12 if not is_header else 12,
                True if is_header or c == 0 or is_amount else False,
                WHITE if is_header else (BLUE_DARK if is_amount else INK),
            )
            fill = NAVY if is_header else (WHITE if r % 2 else SOFT_GREEN)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    add_text(
        s,
        "Payment schedule, validity, and what happens at renewal are written on the quote. This slide is the structure, not the price.",
        0.5,
        6.15,
        12.33,
        0.7,
        14,
        False,
        MUTED,
    )
    footer(s, 14)
    notes(s, "Leave the brackets blank until finance approves a number. Do not invent a discount, a package name, or a market rate.")


def why(prs):
    s = new_slide(prs)
    kicker(s, "Why Elecnovatech", 0.5, 0.28)
    title(s, "The product was shaped on the work it sells.", 0.5, 0.52, 12)
    reasons = [
        ("01", "Our own plants", "The reference installation is Elecnovatech’s block factory and plastic recycling factory. The screens follow that work."),
        ("02", "Local money", "Birr, VAT, withholding, and customer TIN are already part of an order. They are not a translation project."),
        ("03", "Sage stays", "Clients who already close their books in Sage 50 keep Sage. HYBRID is the operating record beside it."),
        ("04", "Jobs are built in", "Owner through driver already have different permissions. A new client maps people onto those jobs, then adjusts them."),
        ("05", "Straight scope", "We separate what is in the product, what is a setting, and what must be quoted. The proposal uses the same three words."),
    ]
    for i, (n, head, body) in enumerate(reasons):
        y = 1.4 + i * 1.08
        add_text(s, n, 0.5, y, 0.7, 0.4, 18, True, GREEN_DARK)
        add_text(s, head, 1.35, y, 3.1, 0.4, 18, True, NAVY)
        add_text(s, body, 4.55, y, 8.2, 0.9, 15, False, SLATE)
    footer(s, 15)
    notes(
        s,
        "Do not cite other customers, uptime, or a payback period. We do not have those as verified sales facts. "
        "Elecnovatech is the company that operates the reference plants and sells the product.",
    )


def close(prs):
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, 0.14, H, GREEN)
    rect(s, 0.14, 0, 0.06, H, BLUE)
    add_text(s, "NEXT CONVERSATION", 0.7, 0.4, 8, 0.3, 13, True, GREEN)
    add_text(s, "See it on their work,\nthen price the written scope.", 0.7, 0.8, 12, 1.5, 32, True, WHITE)
    steps = [
        ("1", "Walk one block size or one recycling day"),
        ("2", "Open one customer who buys on credit"),
        ("3", "Show the Sage comparison, including what the connector still needs"),
        ("4", "Leave with a quote that fills the brackets"),
    ]
    for i, (n, text) in enumerate(steps):
        y = 2.7 + i * 0.72
        rect(s, 0.7, y, 0.42, 0.42, GREEN if i < 3 else BLUE, radius=0.2)
        add_text(s, n, 0.7, y + 0.04, 0.42, 0.34, 14, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, text, 1.3, y + 0.02, 7.5, 0.4, 16, False, RGBColor(0xE4, 0xEC, 0xE6))
    card(s, 9.35, 2.7, 3.4, 3.15, RGBColor(0x16, 0x28, 0x36))
    add_text(s, "ELE CNOVATECH".replace(" ", ""), 9.55, 2.9, 3.05, 0.3, 12, True, GREEN)
    add_text(s, "[Contact name]", 9.55, 3.35, 3.05, 0.35, 16, True, WHITE)
    add_text(s, "[Email]", 9.55, 3.8, 3.05, 0.32, 14, False, RGBColor(0xD0, 0xDE, 0xD4))
    add_text(s, "[Phone]", 9.55, 4.18, 3.05, 0.32, 14, False, RGBColor(0xD0, 0xDE, 0xD4))
    add_text(s, "[Client name]", 9.55, 4.8, 3.05, 0.55, 13, False, RGBColor(0x9A, 0xB4, 0xA0))
    footer(s, 16, dark=True)
    notes(s, "Put the real contact in the card before the meeting. Ask which site they want to see first, and which Sage version they run.")


if __name__ == "__main__":
    build()
