"""Elecnovatech HYBRID product proposal.

Bracketed fields are completed for each client before the document is sent.
Confirmed statements match the product. Proposed items are labeled.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
MARK = ROOT / "assets" / "hybrid-mark.png"
OUT = ROOT / "Elecnovatech-HYBRID-Product-Proposal.docx"

NAVY = RGBColor(0x1A, 0x20, 0x2C)
GREEN = RGBColor(0x2D, 0x5A, 0x27)
GREEN_BRIGHT = RGBColor(0x5C, 0x96, 0x22)
BLUE = RGBColor(0x00, 0x54, 0xA6)
SLATE = RGBColor(0x3D, 0x4A, 0x57)
MUTED = RGBColor(0x5C, 0x6B, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CONTENT_WIDTH = Inches(6.77)


def rgb_hex(color: RGBColor) -> str:
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def set_style_font(style, name, size, bold, color, before=0, after=8):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    style.font.italic = False
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.08
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def shade_cell(cell, fill):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def inches_to_dxa(value):
    return int(value * 1440)


def add_bottom_border(paragraph, color="80C342", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def para(doc, text, size=11, bold=False, italic=False, color=NAVY, space_after=8, space_before=0, align="left"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.alignment = alignment
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def rich_para(doc, parts, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run_font(run, **kwargs)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    if level:
        p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    if p.runs:
        run = p.runs[0]
        run.text = text
        for extra in p.runs[1:]:
            extra.text = ""
    else:
        run = p.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    return p


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    add_bottom_border(p, "80C342", "12")
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def callout(doc, label, text, fill="F3F8EC", label_color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [inches_to_dxa(6.77)])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, 80, 80, 120, 120)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    label_run = p.add_run(label + "   ")
    set_run_font(label_run, size=11, bold=True, color=label_color)
    body = p.add_run(text)
    set_run_font(body, size=11, color=NAVY)
    para(doc, "", size=6, space_after=6)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    dxa = [inches_to_dxa(w) for w in widths]
    set_table_widths(table, dxa)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "1A202C")
        set_cell_margins(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=10, bold=True, color=WHITE)
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            shade_cell(cell, "F6F8F4" if r % 2 else "FFFFFF")
            set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            bold = c == 0
            color = BLUE if value.startswith("[") else NAVY
            set_run_font(run, size=10, bold=bold, color=color)
        prevent_row_split(table.rows[r + 1])
    para(doc, "", size=4, space_after=4)
    return table


def page_break(doc):
    doc.add_page_break()


def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, "Calibri", 11, False, NAVY, 0, 8)
    set_style_font(doc.styles["Heading 1"], "Calibri", 20, True, GREEN, 16, 8)
    set_style_font(doc.styles["Heading 2"], "Calibri", 14, True, BLUE, 14, 4)
    set_style_font(doc.styles["Heading 3"], "Calibri", 12, True, NAVY, 10, 3)
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        set_style_font(style, "Calibri", 11, False, NAVY, 1, 3)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(0.95)
        section.bottom_margin = Inches(0.75)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        section.different_first_page_header_footer = True

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(2)
        r1 = hp.add_run("ELECNOVATECH")
        set_run_font(r1, size=9, bold=True, color=GREEN)
        r2 = hp.add_run("    ·    HYBRID")
        set_run_font(r2, size=9, color=BLUE)
        r3 = hp.add_run("    ·    Product proposal")
        set_run_font(r3, size=9, color=MUTED)
        add_bottom_border(hp, "80C342", "10")

        first_header = section.first_page_header
        first_header.is_linked_to_previous = False
        first_header.paragraphs[0].text = ""

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.space_before = Pt(2)
        add_bottom_border(fp, "E1E8DF", "6")
        fr = fp.add_run("Bracketed fields are completed on the client quote.    ")
        set_run_font(fr, size=8, color=MUTED)
        add_page_field(fp)
        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False
        first_footer.paragraphs[0].text = ""


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8, bold=True, color=GREEN)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_toc(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Open this document in Word and update the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)
    set_run_font(run, size=11, color=SLATE)


def enable_update_fields(doc):
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    run.add_picture(str(MARK), width=Inches(1.15))
    para(doc, "ELECNOVATECH", size=12, bold=True, color=GREEN, space_after=0, space_before=8)
    para(doc, "PRODUCT PROPOSAL", size=11, bold=True, color=BLUE, space_after=8)
    para(doc, "HYBRID", size=36, bold=True, color=NAVY, space_after=2)
    para(
        doc,
        "Operations platform for block factories and plastic recycling plants.",
        size=16,
        color=GREEN,
        space_after=8,
    )
    line = para(doc, "", size=4, space_after=8)
    add_bottom_border(line, "80C342", "16")
    para(
        doc,
        "A reusable product, developed on Elecnovatech’s own block plant and plastic recycling plant, and configured for each new client.",
        size=12,
        color=SLATE,
        space_after=12,
    )
    add_table(
        doc,
        ["Field", "Detail"],
        [
            ("Prepared for", "[Client legal name]"),
            ("Attention", "[Buyer name and role]"),
            ("Prepared by", "Elecnovatech"),
            ("Contact", "[Name]  ·  [Email]  ·  [Phone]"),
            ("Date", "[Date]"),
            ("Reference", "[HYB-____]"),
            ("Valid until", "[Date]"),
            ("Sites in this conversation", "[Block plant / Recycling plant / Both]"),
        ],
        [2.15, 4.62],
    )
    callout(
        doc,
        "HOW TO USE THIS DOCUMENT",
        "The PowerPoint is the meeting. This document is what the buyer keeps. Replace every bracket before it is sent. The words Confirmed, Configurable, and Proposed mean the same thing here as they do on the slides.",
    )
    add_table(
        doc,
        ["In the room", "In this document", "Before it is sent"],
        [
            ("16 slides, one idea each", "The scope a buyer can sign", "Replace every bracket"),
            ("A live look at the screens", "What is confirmed, configured, or still proposed", "Copy the fees onto the slides"),
            ("Their sizes, taxes, and Sage version", "Implementation, support, and acceptance", "Delete any proposed line they did not accept"),
        ],
        [2.25, 2.26, 2.26],
    )


def section_how(doc):
    h1(doc, "1.  How to read this proposal")
    para(
        doc,
        "HYBRID is sold as a configured deployment for one client organization. The first organization it was built for is Elecnovatech’s own operation: a block factory and a plastic recycling factory in the same company. A later client receives the same product, with their own database, their own users, and the settings that match their sizes, taxes, and roles.",
    )
    para(doc, "Three labels are used wherever a claim could be misunderstood.")
    add_table(
        doc,
        ["Label", "Meaning in a sale"],
        [
            ("Confirmed", "A person can do this in the product today, on the screens that are switched on."),
            ("Configurable", "The behavior is already in the product. The client’s values are entered as settings during implementation."),
            ("Proposed", "It is a reasonable next piece of work. It is not included until it is written into the scope and the quote."),
        ],
        [1.6, 5.17],
    )
    para(doc, "Before this document is sent, complete the editing list.")
    for item in [
        "Client legal name, buyer, date, reference, and validity.",
        "Elecnovatech contact name, email, and phone.",
        "Which sites are in scope: block, recycling, or both.",
        "Sage 50 version, or a clear statement that Sage is out of scope.",
        "Annual license, implementation fee, connector fee, customization fee, and monthly support fee.",
        "Payment schedule, timetable, hosting choice, and support hours.",
        "Any proposed item the client has asked to include. If it is not listed in section 11, it is not part of the sale.",
    ]:
        bullet(doc, item)


def section_summary(doc):
    h1(doc, "2.  Executive summary")
    para(
        doc,
        "HYBRID gives a block factory, a plastic recycling plant, or a company that runs both, one operating record. Production, stock, purchasing, sales, credit, banks, expenses, and people sit in the same system. The owner can look at one site or at the company. Finance can compare that record with Sage 50. Floor staff and drivers work from a short set of actions on a phone.",
    )
    para(
        doc,
        "The commercial offer is a license for one organization, an implementation project, an optional Sage 50 connector, only the customization that is written down, and a support retainer after go-live. Amounts are entered in section 13. This document does not publish a price list.",
    )
    add_table(
        doc,
        ["What the client is deciding", "Status"],
        [
            ("Use HYBRID for one organization", "Confirmed product, licensed per the quote"),
            ("Block view, recycling view, and a combined company view", "Confirmed"),
            ("Ethiopian Birr, VAT, withholding, and customer TIN", "Confirmed, with the client’s rates configurable"),
            ("Sage 50 comparison", "Confirmed workspace. The live reader is completed for their Sage version"),
            ("Phone actions for floor, warehouse, and drivers", "Confirmed, in the phone browser"),
            ("Amharic screens, a third site, or a different process", "Proposed, and quoted only if requested"),
        ],
        [4.55, 2.22],
    )
    callout(
        doc,
        "THE DECISION THIS DOCUMENT SUPPORTS",
        "Whether to proceed to a working demonstration on the client’s own sizes or recycling day, and then to a quote that fills the brackets.",
        fill="F0F8FC",
        label_color=BLUE,
    )


def section_problem(doc):
    h1(doc, "3.  The problem the product is sold against")
    para(
        doc,
        "A block plant and a recycling plant generate the same kind of strain, even when the materials are different. Counts start at the machine. Money starts when a customer takes goods on account. The books often already live in Sage 50. Between those three places, the owner is asked to trust numbers that have been copied by hand.",
    )
    add_table(
        doc,
        ["Where the strain shows", "What the buyer usually describes", "What HYBRID is for"],
        [
            ("Production", "Output, damage, and material use are recorded away from the office.", "Batches, daily output, and recycling kilograms in the same system as stock."),
            ("Stock", "Cement, poliche, finished blocks, plastic input, and pellets do not share one balance.", "Live balances, a stock ledger, and physical counts, including from the phone."),
            ("Credit", "A load can leave before anyone can see the real balance.", "Orders carry VAT, discount, withholding, and advances. The amount still owed is calculated."),
            ("The books", "Sage 50 is closed on a different rhythm from the yard.", "A reconciliation workspace, plus a one-way agent from the Sage PC."),
            ("People", "Attendance, leave, and pay sit in another file, with salary visible too widely.", "HR rules, payroll calculation, and a permission that hides salary."),
        ],
        [1.45, 2.7, 2.62],
    )
    para(
        doc,
        "The sales conversation should use the client’s own examples of these strains. HYBRID is not offered with a promised percentage of time saved, waste reduced, or cash collected. Those figures would have to be measured on that client’s data after go-live.",
        italic=True,
        color=SLATE,
    )


def section_product(doc):
    h1(doc, "4.  Product overview")
    h2(doc, "4.1  What the client receives")
    para(
        doc,
        "HYBRID is a web application. People sign in from the office browser. Floor staff and drivers use the same application on a phone browser. The database holds one organization. Inside that organization the product is built for two operating views, Block Factory and Plastic Recycling, and for a combined company view.",
    )
    para(
        doc,
        "A client who runs only one of those businesses works in that view. Renaming the two views, adding a third plant, or covering a process that is neither block-making nor plastic recycling is proposed customization.",
    )
    h2(doc, "4.2  What a working day looks like")
    for line in [
        "The owner opens the company view, then switches to the site that needs attention.",
        "The factory records a batch or a recycling day. Stock balances move with production and with sales.",
        "Sales raises an order for a regular customer or a walking customer. VAT, discount, withholding, and any advance sit on that order.",
        "The driver confirms pickup or delivery with a photo, a signature, and a GPS point. If the signal drops, the phone keeps the action and sends it later.",
        "Finance reviews collections, receivables, expenses, and the comparison with Sage 50.",
        "HR records attendance and leave against the company’s own calendar and rates, and runs payroll inside those rules.",
    ]:
        bullet(doc, line)
    h2(doc, "4.3  What the product is deliberate about")
    para(
        doc,
        "Numbers that matter, including yield, capacity, waste, and credit scores, are calculated in the application from the client’s own records. When an AI service key is configured, the assistant can answer in plain language. It is instructed to take figures from those calculations. Plain-language answers are absent when no key is configured. The calculations remain.",
    )
    callout(
        doc,
        "CONFIRMED BOUNDARY",
        "The screens are in English. Amounts are in Ethiopian Birr. An Amharic interface is a proposed customization, not part of the current product.",
    )


def section_market(doc):
    h1(doc, "5.  Who it is for")
    para(doc, "HYBRID is a fit for a business that recognizes itself in at least one of these descriptions.")
    add_table(
        doc,
        ["Buyer", "The work HYBRID already understands"],
        [
            ("Block manufacturer", "Finished sizes, cement bags, plates, poliche loads, pallets, and sales of those blocks."),
            ("Plastic recycling plant", "Kilograms received, kilograms processed, waste, and sales of the output."),
            ("A company with both", "Shared customers, suppliers, and banks, with a separate operating view for each site."),
            ("The management team", "An owner, a finance lead, a factory manager, warehouse, sales, HR, floor staff, and drivers."),
        ],
        [2.1, 4.67],
    )
    h2(doc, "5.1  A strong first conversation")
    para(doc, "The first meeting is strongest when the buyer can answer yes to most of these.")
    for line in [
        "They sell on account, and they need to know the balance before the next load.",
        "They already use Sage 50, or they want the operating record kept ready for that comparison.",
        "They charge VAT or withholding, and customers have a TIN.",
        "More than one job needs access, and salary or credit should not be visible to everyone.",
        "Someone on the floor or on a truck needs to record a count or a delivery without sitting at a desk.",
    ]:
        bullet(doc, line)
    h2(doc, "5.2  Where a sale should slow down")
    para(
        doc,
        "If the client’s process is not block production and not plastic recycling, HYBRID is the wrong product to promise as it stands. The right response is a written customization scope, or a decision not to proceed. The same applies to a request for a shared website on which unrelated companies would log in together. The current product is one organization per deployment.",
    )


def section_benefits(doc):
    h1(doc, "6.  Benefits by role")
    para(
        doc,
        "Benefits below are the practical consequence of the confirmed screens. They are not measured outcomes. A proposal should not convert them into a payback figure.",
    )
    add_table(
        doc,
        ["Role", "What they can do", "Why it matters on a Tuesday"],
        [
            ("Owner", "See the company, the block plant, or the recycling plant.", "The question “where do we stand?” has one place to start."),
            ("Finance manager", "Statements, aging, payments, expenses, banks, and the Sage comparison.", "The factory amount and the ledger amount can be looked at together."),
            ("Factory manager", "Batches, targets, yield, damage, and capacity from materials on hand.", "The next run can be planned from stock that is already in the system."),
            ("Warehouse", "Balances, physical counts, and movements from the phone.", "A count does not wait for someone to retype a sheet."),
            ("Sales", "Customers, orders, and the credit position.", "A pending order shows what is still owed, with withholding already applied."),
            ("Driver", "Delivery proof only: photo, signature, GPS.", "The office can see that the load was closed, including when the phone was offline."),
            ("HR manager", "Attendance, leave, corrections, and payroll from company rules.", "Pay is calculated from the rules and the attendance, and salary stays behind a permission."),
            ("Floor employee", "Production count and stock count.", "The job on the phone matches the job in the yard."),
        ],
        [1.45, 2.7, 2.62],
    )


def section_capabilities(doc):
    h1(doc, "7.  Confirmed capabilities")
    para(
        doc,
        "This section is the catalog a buyer, and later an implementer, can check against the product. Where a screen exists in the code and is switched off in the current menu, it is listed as proposed in section 11.",
    )

    h2(doc, "7.1  Company and sites")
    bullet(doc, "Confirmed. One organization in a deployment. Users belong to that organization.")
    bullet(doc, "Confirmed. Three views: the whole company, the block factory, and the plastic recycling plant.")
    bullet(doc, "Confirmed. Currency is Ethiopian Birr. Number formatting follows an Ethiopia locale.")
    bullet(doc, "Configurable. Which view a person works in, and which permissions that person has.")

    h2(doc, "7.2  Block production")
    bullet(doc, "Confirmed. Production batches for 10 cm, 15 cm, 20 cm, and 24 cm blocks.")
    bullet(doc, "Confirmed. Standard conversions among pieces, plates, pallets, cement bags, and poliche loads. The default cement bag is 45 kg, and that weight is a setting.")
    bullet(doc, "Confirmed. Counts can be recorded at batch count, after drying, and after curing, together with daily output.")
    bullet(doc, "Confirmed. Yield and damage are calculated from the counts entered.")
    bullet(doc, "Confirmed. Daily targets can be set. A capacity view estimates how many blocks the current cement and poliche can still make, and which material runs out first.")
    bullet(doc, "Confirmed. Planning figures drawn from daily history carry a confidence level. The level reflects how much history exists. It is not a guaranteed forecast accuracy.")
    bullet(doc, "Confirmed. An in-app calculator converts block quantities and estimates materials, waste, and related amounts.")
    bullet(doc, "Configurable. Yield formulas, bag weight, and daily targets.")

    h2(doc, "7.3  Plastic recycling")
    bullet(doc, "Confirmed. A recycling entry records kilograms received, kilograms processed, and waste.")
    bullet(doc, "Confirmed. Yield is calculated from received and processed quantities.")
    bullet(doc, "Confirmed. The recycling view shows raw-material stock, finished stock, and output for today, the week, and the month.")
    bullet(doc, "Confirmed. A recycling order has its own stages: pending, accepted, in production, ready, delivered, and cancelled.")
    bullet(doc, "Confirmed. Recycling output, yield, and waste shown in planning come from those records. The block-size learning model is not applied to recycling.")

    h2(doc, "7.4  Inventory, purchasing, and suppliers")
    bullet(doc, "Confirmed. Live on-hand balances and a stock ledger.")
    bullet(doc, "Confirmed. Physical stock counts, including a phone action.")
    bullet(doc, "Confirmed. Low-stock markers on items that have a threshold.")
    bullet(doc, "Confirmed. A catalog of item types, categories, and units, edited in administration.")
    bullet(doc, "Confirmed. Purchasing records, supplier accounts, and supplier balances.")
    bullet(doc, "Confirmed. Phone action to log a stock movement. The separate movements tab on the inventory screen is switched off today. Switching it on is proposed.")
    bullet(doc, "Proposed, already built and hidden. Bill of materials, locations, and lots. Offer these only as a scoped switch-on.")

    h2(doc, "7.5  Sales, customers, and delivery")
    bullet(doc, "Confirmed. Customers are shared across both sites. A customer is regular or walking, and can store a TIN.")
    bullet(doc, "Confirmed. Order status is pending pickup, picked, or cancelled. Recycling orders also carry the stage list above.")
    bullet(doc, "Confirmed. An order can carry VAT, a discount, withholding, and an advance or prepaid amount.")
    bullet(doc, "Confirmed. Payment method on an order is cash or bank. An order can be marked as requiring an invoice.")
    bullet(doc, "Confirmed. The remaining credit amount is calculated so that withholding and store credit are already reflected. Staff are not expected to subtract those again by hand.")
    bullet(doc, "Confirmed. Delivery confirmation captures a photo, a signature, and a GPS point.")

    h2(doc, "7.6  Credit")
    bullet(doc, "Confirmed. Customer balances, payments, and a credit screen.")
    bullet(doc, "Confirmed. A score from 0 to 100, with a risk tier and a limit, calculated from that customer’s own history.")
    bullet(doc, "Confirmed. Separate scores for the block plant, the recycling plant, and the company combined.")
    bullet(doc, "Confirmed boundary. The score is an internal figure for the client’s own credit decisions. It is not a bank score and it does not approve a loan.")

    h2(doc, "7.7  Finance, tax, banks, and expenses")
    bullet(doc, "Confirmed. Period views for revenue, collections, outstanding receivables, and cash signals.")
    bullet(doc, "Confirmed. Customer statements, receivables aging (0–30, 31–60, 61–90, and over 90 days), payment report, revenue summary, expense report, profitability, and bank transactions.")
    bullet(doc, "Confirmed. Bank accounts, including a place for an account number.")
    bullet(doc, "Confirmed. Operating expenses, with an analytics view.")
    bullet(doc, "Configurable. VAT rate and withholding rate. Withholding applies to new orders across the company once the rate is set.")
    bullet(doc, "Confirmed boundary. Payroll amounts and order taxes are calculated in HYBRID. The product does not file tax returns.")

    h2(doc, "7.8  Sage 50")
    para(
        doc,
        "The reference operation compares factory amounts with Sage 50. HYBRID includes that comparison as a reconciliation workspace. Statuses are matched, variance, pending, and approved. A Windows agent installed on the Sage computer sends data one way to HYBRID, on a timer or when someone presses Sync Now. The agent authenticates with a key and a secret. It does not require an inbound connection to the client’s network.",
    )
    callout(
        doc,
        "CONFIRMED LIMIT",
        "The agent that ships reads an export file. A live reader for the client’s exact Sage 50 version is completed during implementation, and only if Sage is in the quote. HYBRID does not replace Sage 50.",
        fill="FBF6EA",
        label_color=RGBColor(0x8A, 0x6A, 0x1A),
    )

    h2(doc, "7.9  Human resources")
    bullet(doc, "Confirmed. Employee records.")
    bullet(doc, "Confirmed. Attendance against a work schedule, including late arrival and early departure from the company’s clock rules.")
    bullet(doc, "Confirmed. Leave types, accrual, holidays, probation handling, approval, and attendance corrections.")
    bullet(doc, "Confirmed. Leave compensation can be switched on, with the company’s own daily-rate rule.")
    bullet(doc, "Confirmed. Payroll is calculated from salary, attendance, overtime, leave, and the configured payroll tax rate. Payslips are produced in the product.")
    bullet(doc, "Confirmed. Viewing salary is a separate permission from viewing the rest of HR.")
    bullet(doc, "Configurable. Work hours, grace minutes, leave types, holidays, shifts, payroll tax rate, and whether leave converts to pay.")

    h2(doc, "7.10  Reports, documents, notifications, and audit")
    bullet(doc, "Confirmed. On-demand export to PDF and Excel for the report list in Appendix C, with filters such as date range.")
    bullet(doc, "Confirmed. Report schedule records can be stored and paused. Automatic email of those schedules is not a verified mail connection, and is proposed if the client wants it.")
    bullet(doc, "Confirmed. Document register with tags, restricted files, and a link to the related record.")
    bullet(doc, "Confirmed. In-app notifications, including alerts raised from the client’s own data.")
    bullet(doc, "Confirmed. An audit log of changes to key records, with its own export.")

    h2(doc, "7.11  Roles and access")
    para(doc, "Confirmed roles shipped with the product:")
    add_table(
        doc,
        ["Role", "The job it is meant for"],
        [
            ("Owner / CEO", "The whole company."),
            ("Finance manager", "Finance, sales read, reports, documents, audit, credit view."),
            ("Factory manager", "Production, inventory, expenses, reports, capacity."),
            ("Production supervisor", "Production entry and verification, capacity."),
            ("Warehouse officer", "Stock, counts, and production read."),
            ("Sales manager", "Orders, customers, delivery, credit view."),
            ("HR manager", "People, leave approval, payroll, salary, policies."),
            ("Floor employee", "Production counts and stock counts."),
            ("Driver", "Delivery confirmation."),
        ],
        [2.2, 4.57],
    )
    para(doc, "Configurable. Permissions are stored as data. An administrator can adjust what a role is allowed to do. Every sensitive action is checked on the server, and the database uses row-level security.")

    h2(doc, "7.12  Phone")
    bullet(doc, "Confirmed. The phone experience is the web application, laid out for a small screen. It is not a separate listing in an app store.")
    bullet(doc, "Confirmed. Floor and production supervisors see production count and stock count.")
    bullet(doc, "Confirmed. Warehouse sees stock count and inventory movement.")
    bullet(doc, "Confirmed. Drivers see delivery confirmation only.")
    bullet(doc, "Confirmed. Factory managers see all four actions.")
    bullet(doc, "Confirmed. Actions can be queued while the phone is offline and sent when the connection returns.")

    h2(doc, "7.13  Planning assistant")
    bullet(doc, "Confirmed. The assistant can be asked about inventory, capacity, production history, a hypothetical block run, batch status, finance for a period, a customer’s credit, open orders, and supplier balances.")
    bullet(doc, "Confirmed. Answers are limited by the signed-in person’s role and by the site currently in view.")
    bullet(doc, "Confirmed. Plain language requires an xAI or OpenAI key, supplied for that deployment. Without a key, the calculated screens remain.")
    bullet(doc, "Confirmed boundary. The assistant is not a source of numbers. The numbers are computed in HYBRID.")


def section_ethiopia(doc):
    h1(doc, "8.  Value for block and recycling businesses in Ethiopia")
    para(
        doc,
        "HYBRID is not a generic operations package with Ethiopian labels added afterward. The reference operation is in Ethiopia, the currency is Birr, and the commercial documents a plant actually issues are represented on the order.",
    )
    add_table(
        doc,
        ["Local reality", "How the product meets it", "Status"],
        [
            ("Prices and books in Birr", "The application currency is ETB.", "Confirmed"),
            ("VAT on sales", "A company VAT setting, applied on orders.", "Configurable"),
            ("Withholding", "A company withholding setting on new orders, already reflected in the amount owed.", "Configurable"),
            ("Customer TIN", "Stored on the customer.", "Confirmed"),
            ("Cash and bank collections", "Order payment method is cash or bank. Bank accounts are first-class records.", "Confirmed"),
            ("Sage 50 beside the yard", "Reconciliation workspace and a Windows agent.", "Confirmed path; reader is per version"),
            ("Block sizes and poliche", "10, 15, 20, and 24 cm, with poliche loads in the formula.", "Confirmed; formulas configurable"),
            ("Plastic in kilograms", "Received, processed, and waste.", "Confirmed"),
            ("English screens", "The interface language is English.", "Confirmed"),
            ("Amharic screens", "Not in the product.", "Proposed"),
        ],
        [2.05, 3.15, 1.57],
    )
    para(
        doc,
        "This is the practical value in a buying meeting. The owner does not have to translate a foreign factory system into bags, poliche, withholding, and two sites. Those are already the objects on the screen. What still has to be translated is the client’s own rates, sizes, and people’s names, and that work is the implementation.",
    )


def section_origin(doc):
    h1(doc, "9.  Where the product comes from")
    para(
        doc,
        "Elecnovatech operates a block factory and a plastic recycling factory. HYBRID was built as the operating system for that company. The product name in the software is HYBRID. The two sites in the software are that block factory and that recycling plant.",
    )
    para(
        doc,
        "The commercial step is to license the same product to another business that does this work. Each new client is a new organization, with its own deployment, its own users, and its own settings. Elecnovatech remains the company that sells, implements, and supports it.",
    )
    callout(
        doc,
        "NO THIRD-PARTY REFERENCES IN THIS DOCUMENT",
        "There are no customer testimonials, no count of other clients, and no performance percentage in this proposal. The reference is Elecnovatech’s own operation. If a later client agrees to be named, that agreement is added in writing. It is not assumed here.",
        fill="F0F8FC",
        label_color=BLUE,
    )
    h2(doc, "9.1  Why that origin is an advantage")
    for line in [
        "The units on the screen were required by a plant that actually makes blocks and a plant that actually recycles plastic.",
        "Roles already distinguish an owner, finance, the factory, the warehouse, sales, HR, the floor, and a driver.",
        "Credit, withholding, and Sage 50 are in the product because the reference operation needed them, not because they were added as a brochure line.",
        "A new client can be shown the real screens, then told plainly which settings will change and which requests are extra scope.",
    ]:
        bullet(doc, line)


def section_implementation(doc):
    h1(doc, "10.  Implementation")
    para(
        doc,
        "Implementation is a project with a written timetable, shown in the quote as [Agreed timetable]. The steps below are the method. They are not a promise that every client finishes in the same number of weeks.",
    )
    add_table(
        doc,
        ["Step", "Elecnovatech", "The client"],
        [
            ("1. Discover", "Facilitates the workshop and writes the scope.", "Brings sizes, tax rates, banks, roles, and the Sage version if any."),
            ("2. Deploy", "Creates the database, sign-in, and the first administrator.", "Names the person who will administer users."),
            ("3. Configure", "Enters yields, catalog, tax, HR rules, roles, and banks.", "Confirms the values before go-live."),
            ("4. Open the books", "Supports entry of opening stock, customers, suppliers, and employees.", "Provides the opening lists and checks them."),
            ("5. Connect Sage", "Completes the reader for the agreed Sage 50 version and installs the agent.", "Provides the Windows PC where Sage runs, and a person who can approve it. Skipped if Sage is out of scope."),
            ("6. Train", "Runs a session for each role in scope.", "Releases those people for the session."),
            ("7. Go live", "Stays available for the hypercare window in the quote.", "Runs live orders, counts, and payments in HYBRID from the agreed date."),
        ],
        [1.35, 2.71, 2.71],
    )
    h2(doc, "10.1  What “done” means")
    para(doc, "Go-live is complete when all of the following are true.")
    for line in [
        "The agreed users can sign in and see only their role.",
        "The agreed sites are in use, with the confirmed formulas and tax rates.",
        "Opening stock and the opening customer and supplier lists have been checked by the client.",
        "At least one order path has been practiced: create, pick or stage, and, if drivers are in scope, delivery proof.",
        "If Sage is in scope, a sync has been run and a variance can be shown in the reconciliation workspace.",
        "The client has the quote’s support contact and the hypercare end date.",
    ]:
        bullet(doc, line)
    h2(doc, "10.2  Data at the start")
    para(
        doc,
        "Opening balances and master data are entered during implementation. A spreadsheet import is not a confirmed feature. If the client wants a one-time load from their existing sheets, that load is written into the customization scope and priced there.",
    )
    h2(doc, "10.3  Hosting")
    para(
        doc,
        "The application runs against a Supabase project: authentication, PostgreSQL, file storage, and row-level security. The quote selects one of these: [Client-managed project] or [Elecnovatech-managed hosting, at the fee stated in the quote]. The client’s operating records belong to the client. A copy of how to export them is part of the handover.",
    )


def section_custom(doc):
    h1(doc, "11.  Configuration and customization")
    h2(doc, "11.1  Included as configuration")
    para(doc, "These changes are part of implementation. They do not need a separate customization fee.")
    for line in [
        "Block yield formulas, cement bag weight, and daily targets.",
        "VAT, withholding, and the payroll tax rate the company chooses to apply inside HYBRID.",
        "Work hours, grace time, leave types, holidays, shifts, and leave-to-pay rules.",
        "Inventory types, categories, and units.",
        "Bank accounts, customer and supplier records, and employee records.",
        "Role permissions, within the resources the product already checks.",
        "Which of the two sites are active for this client.",
    ]:
        bullet(doc, line)
    h2(doc, "11.2  Quoted only when written down")
    add_table(
        doc,
        ["Request", "Why it is not already included"],
        [
            ("Sage 50 reader for a named version", "The shipped agent reads a file. The live reader depends on that version."),
            ("Amharic, or any other interface language", "The screens are English."),
            ("Switch on movements, bill of materials, locations, or lots", "The code exists. Those inventory tabs are hidden."),
            ("A third site, or a site that is not block or recycling", "The product is modeled as two sites plus a company view."),
            ("Rename the sites to the client’s plant names", "The current labels are Block Factory and Plastic Recycling."),
            ("Email delivery of scheduled reports", "Schedules can be stored. A mail integration is not verified."),
            ("A one-time import from spreadsheets", "Opening data is entered. Import is extra work."),
            ("A native app-store application", "The phone tool is the website."),
            ("A second authentication step", "Sign-in is the account password. A role flag for multi-factor exists. A working second step does not."),
            ("One website shared by unrelated companies", "Each client is a separate deployment."),
            ("Any workflow the discovery workshop cannot map onto the screens", "It is new product work, scoped and priced on its own."),
        ],
        [2.85, 3.92],
    )
    callout(
        doc,
        "RULE FOR REPEATED SALES",
        "If it is not in section 7 or section 11.1, it goes on the customization list with a fee, or it stays out of the sale. Verbal agreement in a meeting is not scope.",
        fill="FBF6EA",
        label_color=RGBColor(0x8A, 0x6A, 0x1A),
    )


def section_support(doc):
    h1(doc, "12.  Support and maintenance")
    para(
        doc,
        "Support terms below are a proposed commercial package. They become binding when the quote names the hours, the fee, and the hypercare length. They are not a description of an existing service-level contract.",
        italic=True,
        color=SLATE,
    )
    add_table(
        doc,
        ["Period", "What it covers", "Commercial treatment"],
        [
            ("Implementation", "The steps in section 10, through go-live.", "Included in [Project fee]."),
            ("Hypercare", "Questions and defect correction while the client starts live use.", "[Number] weeks after go-live, included unless the quote says otherwise."),
            ("Retainer", "Questions on confirmed features, defects in the licensed product, and small setting changes listed in the quote.", "[Monthly fee], starting [date]."),
            ("Out of retainer", "New workflows, new sites, Sage version changes, new languages, and training for people hired later.", "Quoted separately."),
        ],
        [1.45, 3.15, 2.17],
    )
    h2(doc, "12.1  Response, once the quote sets it")
    para(doc, "Write the actual targets into the quote. Suggested blanks:")
    bullet(doc, "Critical: the client cannot record production, orders, or payments. Response within [hours] on [days].")
    bullet(doc, "Other: questions and small setting changes. Response within [business days].")
    bullet(doc, "Hours of cover: [for example, weekdays 08:30–17:30, East Africa Time].")
    bullet(doc, "Channel: [phone / email / named chat]. Contact: [name].")
    h2(doc, "12.2  What a defect is")
    para(
        doc,
        "A defect is a confirmed feature behaving differently from this document and from the screens. A request to behave differently is customization. That distinction protects both the client and the next sale.",
    )
    h2(doc, "12.3  Product updates")
    para(
        doc,
        "Corrections to the licensed product during the retainer are included. A new capability the client wants, even if Elecnovatech later builds it for the product in general, is scheduled with the client and may be a new quote. The quote states the term of the license: [term, for example one year], and what renewal costs: [renewal fee or “same as the annual fee”].",
    )


def section_pricing(doc):
    h1(doc, "13.  Pricing approach")
    para(
        doc,
        "HYBRID is priced per organization, not per a public rate card. The structure is fixed so that every proposal looks the same. The amounts are not. Enter them here, and use the same amounts on the presentation’s commercial slide.",
    )
    add_table(
        doc,
        ["Item", "Basis", "Amount"],
        [
            ("License", "Right to use HYBRID for one organization for [term].", "[Annual fee]"),
            ("Implementation", "Section 10, including training and the hypercare window.", "[Project fee]"),
            ("Sage connector", "Live reader for [Sage 50 version], plus installation on the client PC. Use “Not required” if they do not use Sage.", "[Fee or Not required]"),
            ("Customization", "Only the lines copied from section 11.2 that this client accepted.", "[Scoped fee or None]"),
            ("Support retainer", "Section 12, after hypercare.", "[Monthly fee]"),
            ("Hosting", "Elecnovatech-managed Supabase and application hosting, or client-managed if the quote says so.", "[Fee or Client-managed]"),
        ],
        [1.7, 3.35, 1.72],
    )
    h2(doc, "13.1  Suggested payment schedule")
    para(doc, "Edit this schedule. It is a suggestion for the quote, not a term already used with other buyers.")
    bullet(doc, "[Percent]% on written acceptance of this proposal.")
    bullet(doc, "[Percent]% when training starts.")
    bullet(doc, "[Percent]% on the go-live date in section 10.")
    bullet(doc, "The annual license is invoiced [at acceptance / at go-live / annually in advance].")
    bullet(doc, "The retainer is invoiced [monthly in advance], starting [date].")
    h2(doc, "13.2  Currency and validity")
    bullet(doc, "Quote currency: [ETB].")
    bullet(doc, "This proposal is valid until [date]. After that date the amounts are confirmed again.")
    bullet(doc, "Taxes on Elecnovatech’s invoice: [state whether VAT is added].")
    h2(doc, "13.3  What the price is not")
    for line in [
        "It is not a discount against an invented list price.",
        "It does not include proposed items that were discussed and then left off the customization list.",
        "It does not include the client’s Sage 50 licenses, computers, phones, or internet.",
        "It does not include an AI service subscription. If the client wants plain-language answers, the key and its cost are theirs, noted here: [client provides the key / Elecnovatech provides it inside the fee].",
    ]:
        bullet(doc, line)


def section_why(doc):
    h1(doc, "14.  Reasons to choose Elecnovatech")
    reasons = [
        (
            "The work is already the product.",
            "Elecnovatech runs the block plant and the recycling plant the software was built around. A buyer is looking at a system shaped by that operation, then configured for theirs.",
        ),
        (
            "Local commercial documents are native.",
            "Birr, VAT, withholding, TIN, cash, and bank are part of an order. The implementation sets the client’s rates. It does not invent the fields.",
        ),
        (
            "Sage 50 can stay the accounting system.",
            "Clients who already close their books in Sage 50 keep Sage. HYBRID is the operating record and the comparison. The connector is honest about needing their version.",
        ),
        (
            "The jobs already exist.",
            "From owner to driver, the permissions match a factory. The client maps people onto those jobs and then adjusts the permissions that should differ.",
        ),
        (
            "Scope stays written.",
            "Confirmed, configurable, and proposed are the same three words in the meeting and in this document. A buyer can see what they are paying for, and what they are choosing not to buy yet.",
        ),
        (
            "The client’s data stays the client’s.",
            "Each sale is a separate organization. Records are not mixed with Elecnovatech’s plants or with another buyer’s plants.",
        ),
    ]
    for title, body in reasons:
        h3(doc, title)
        para(doc, body)


def section_limits(doc):
    h1(doc, "15.  Statements this proposal will not make")
    para(doc, "Sales staff should treat this list as part of the offer. Saying any of these in a meeting is a defect in the sale.")
    for line in [
        "A percentage improvement in waste, cash collection, or staff time.",
        "A testimonial, a logo, or a count of other HYBRID clients.",
        "That Sage 50 connects on the first day for every version.",
        "That the phone tool is installed from an app store.",
        "That the screens are available in Amharic.",
        "That HYBRID files VAT, withholding, or payroll tax with the authority.",
        "That the credit score is accepted by a bank.",
        "That scheduled reports are emailed today.",
        "That bill of materials, warehouse locations, and lots are on the current inventory menu.",
        "That unrelated companies share one HYBRID login.",
        "A fixed go-live date before the discovery workshop.",
    ]:
        bullet(doc, line)


def section_demo(doc):
    h1(doc, "16.  Suggested demonstration")
    para(
        doc,
        "A demonstration uses the real product. It does not use invented sample results as if they were the client’s. A short path is enough for a first meeting. A second meeting can go into the client’s sizes.",
    )
    add_table(
        doc,
        ["Step", "What to open", "What to say"],
        [
            ("1", "Company, block, and recycling views", "These are the two sites and the company. A single-site client uses one of them."),
            ("2", "A block batch and the calculator", "Sizes, cement, poliche, yield, and damage. Formulas change in settings."),
            ("3", "Recycling stock", "Kilograms in, kilograms out, waste, and the order stages."),
            ("4", "One customer and the credit screen", "Balance, withholding already applied, and a score from their own payments."),
            ("5", "Finance reconciliation", "How a variance is reviewed. Then explain the Sage agent exactly as section 7.8 does."),
            ("6", "The phone, as a driver or a floor user", "Few actions, and a queue when the signal drops."),
        ],
        [0.7, 2.55, 3.52],
    )
    para(doc, "Close the demonstration by writing down anything the client asked for that is not on this path. That list is section 11.2 for their quote.")


def section_next(doc):
    h1(doc, "17.  Next step and acceptance")
    para(doc, "The proposed next step is:")
    bullet(doc, "A demonstration on [date] at [place or video], covering the path in section 16.")
    bullet(doc, "A completed quote, using section 13, sent by [date].")
    bullet(doc, "A decision by the client before the validity date.")
    para(
        doc,
        "Acceptance below authorizes Elecnovatech to start discovery and deployment for the sites named, at the amounts named. It does not authorize work that is still marked proposed and is not copied into the customization lines.",
    )
    h2(doc, "17.1  Scope being accepted")
    add_table(
        doc,
        ["Item", "This client"],
        [
            ("Sites", "[Block / Recycling / Both]"),
            ("Sage 50", "[Version, or Not in scope]"),
            ("Customization lines", "[None, or paste the accepted lines]"),
            ("License", "[Annual fee] for [term]"),
            ("Implementation", "[Project fee]"),
            ("Sage connector", "[Fee or Not required]"),
            ("Customization fee", "[Fee or None]"),
            ("Support retainer", "[Monthly fee] after [hypercare length]"),
            ("Hosting", "[Client-managed or Elecnovatech-managed]"),
            ("Timetable", "[Agreed timetable]"),
            ("Quote currency", "[ETB]"),
        ],
        [2.3, 4.47],
    )
    h2(doc, "17.2  Signatures")
    para(doc, "For the client")
    for line in [
        "Legal name: [Client legal name]",
        "Name: ________________________________",
        "Role: ________________________________",
        "Signature: ________________________________",
        "Date: ________________________________",
    ]:
        para(doc, line, space_after=4)
    para(doc, "For Elecnovatech", space_before=10)
    for line in [
        "Name: ________________________________",
        "Role: ________________________________",
        "Signature: ________________________________",
        "Date: ________________________________",
    ]:
        para(doc, line, space_after=4)
    para(
        doc,
        "Contracting entity, if it differs from the brand name Elecnovatech: [Legal entity name].",
        italic=True,
        color=SLATE,
        space_before=8,
    )


def appendix_tech(doc):
    h1(doc, "Appendix A.  Technical summary for the client’s IT contact")
    para(
        doc,
        "This appendix is descriptive. It is not a security certification, an uptime commitment, or a penetration-test report. Those would be separate documents if a client requires them.",
    )
    add_table(
        doc,
        ["Topic", "What is in place"],
        [
            ("Application", "A web application used from a desktop or phone browser."),
            ("Identity", "Sign-in through Supabase Auth. Access is then limited by role."),
            ("Data", "PostgreSQL, with row-level security. Files such as documents and delivery photos use storage buckets."),
            ("Permissions", "Checked on the server for each action, and editable by an administrator."),
            ("Audit", "An audit log on key record changes."),
            ("Sage agent", "A program on the Windows PC that hosts Sage 50. It sends data outward. It authenticates with a key and a secret."),
            ("AI", "Optional. An xAI or OpenAI key enables plain-language answers. Calculations do not depend on that key."),
            ("Reports", "PDF and Excel generated from the client’s data at the time of export."),
            ("Phone offline", "A queue on the device for the floor actions, sent when the device is online again."),
            ("Isolation", "One organization per deployment. The running configuration is not a shared portal for many companies."),
        ],
        [1.55, 5.22],
    )


def appendix_discovery(doc):
    h1(doc, "Appendix B.  Discovery questions")
    para(doc, "Use this list in the first workshop. The answers become the configuration, or a line in section 11.2.")
    groups = [
        (
            "Sites and products",
            [
                "Block plant, recycling plant, or both?",
                "Which block sizes are actually made? Are any outside 10, 15, 20, and 24 cm?",
                "What is a cement bag in kilograms, and do they use poliche as a unit?",
                "For recycling, what do they call the input and the output, and is the unit the kilogram?",
            ],
        ),
        (
            "Money",
            [
                "VAT rate, and is it on every order?",
                "Withholding rate, and when does it apply?",
                "Do they sell for cash, bank, and on credit?",
                "Which banks receive customer payments?",
                "Do they use Sage 50? Which version, on which Windows PC, and who administers it?",
            ],
        ),
        (
            "People",
            [
                "Who is the owner, the finance lead, the factory manager, and the HR lead?",
                "Who may see salary? Who may approve leave?",
                "Do drivers need delivery proof? Do floor staff have phones with a browser?",
                "What are the working days, the shift times, and the grace period for lateness?",
                "Which leave types and public holidays apply?",
            ],
        ),
        (
            "Start",
            [
                "Who will check opening stock, customers, and suppliers?",
                "Is a one-time spreadsheet load required, or will data be entered?",
                "Who will administer new users after go-live?",
                "Preferred language of the screens. If Amharic is required, it is a customization line.",
            ],
        ),
    ]
    for title, questions in groups:
        h2(doc, title)
        for question in questions:
            bullet(doc, question + "  [Answer]")


def appendix_reports(doc):
    h1(doc, "Appendix C.  Reports included")
    para(doc, "Each of these can be exported to PDF or Excel. The default format is the one the product suggests. Either format can be chosen at export.")
    add_table(
        doc,
        ["Report", "What it contains"],
        [
            ("Production summary", "Batch status and production counts, with date, factory, and checkpoint filters."),
            ("Finance reconciliation", "Operational amounts against Sage amounts, with variance and status."),
            ("Customer statements", "Balances, collections, and order activity for the period."),
            ("Accounts receivable aging", "Outstanding credit in 0–30, 31–60, 61–90, and 90+ days."),
            ("Payment report", "Customer payments with method, reference, and totals."),
            ("Revenue summary", "Picked-order revenue, with payment and credit breakdown."),
            ("Expense report", "Sage expense transactions grouped by factory."),
            ("Profitability", "Revenue minus expenses, with margin for the period."),
            ("Bank transactions", "Customer bank payments and Sage bank entries."),
            ("Finance dashboard summary", "Finance indicators, customer balances, alerts, and the period trend."),
            ("Inventory valuation", "Stock levels, unit costs, and value, with category and health filters."),
            ("Executive dashboard export", "A cross-module summary for a leadership review."),
            ("HR attendance", "Daily attendance, clock times, and status by department."),
            ("Sales and orders pipeline", "Orders with customer, value, delivery dates, and risk flags."),
            ("Customer account summary", "Balances, order activity, and pipeline exposure."),
            ("Audit activity log", "Audit trail by entity, action, and date."),
            ("Document register", "Documents with tags, restrictions, and linked records."),
        ],
        [2.35, 4.42],
    )


def appendix_roles(doc):
    h1(doc, "Appendix D.  Editing checklist for the next sale")
    para(doc, "Copy this proposal for the next client. Then do only this.")
    for line in [
        "Save a new file named with the client and the reference number.",
        "Replace the cover table.",
        "Set sites, Sage version, and customization lines in section 17.1.",
        "Fill section 13 with approved amounts. Copy those amounts onto the presentation.",
        "Delete any proposed item this client did not accept, so the document they sign matches the work.",
        "Read section 15 once more before the meeting.",
        "Put the presenter’s name and phone on the last slide.",
    ]:
        bullet(doc, line)
    para(
        doc,
        "The presentation file that matches this proposal is Elecnovatech-HYBRID-Sales-Presentation.pptx. The two documents should be sent together.",
        space_before=6,
    )


def build():
    doc = Document()
    setup_styles(doc)
    core = doc.core_properties
    core.author = "Elecnovatech"
    core.title = "HYBRID — Product proposal"
    core.subject = "Operations platform for block factories and plastic recycling plants"
    core.category = "Sales proposal"
    core.comments = "Editable proposal. Complete every bracket before sending. Do not add unverified metrics or features."

    cover(doc)
    page_break(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Contents")
    set_run_font(run, size=20, bold=True, color=GREEN)
    add_bottom_border(p, "80C342", "12")
    para(doc, "Update this list in Word if a section is added or renamed.", italic=True, color=MUTED)
    add_toc(doc)
    page_break(doc)

    section_how(doc)
    section_summary(doc)
    section_problem(doc)
    section_product(doc)
    section_market(doc)
    section_benefits(doc)
    section_capabilities(doc)
    section_ethiopia(doc)
    section_origin(doc)
    section_implementation(doc)
    section_custom(doc)
    section_support(doc)
    section_pricing(doc)
    section_why(doc)
    section_limits(doc)
    section_demo(doc)
    section_next(doc)
    appendix_tech(doc)
    appendix_discovery(doc)
    appendix_reports(doc)
    appendix_roles(doc)

    enable_update_fields(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
